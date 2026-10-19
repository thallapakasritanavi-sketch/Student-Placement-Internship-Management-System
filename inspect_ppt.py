from pptx import Presentation

ppt_path = r'C:\Users\thall\OneDrive\Desktop\Professional Training PPT edited.pptx'
prs = Presentation(ppt_path)

print(f'Total slides: {len(prs.slides)}')
for idx, slide in enumerate(prs.slides):
    print(f'\n================= SLIDE {idx+1} =================')
    for s in slide.shapes:
        txt = ""
        if s.has_text_frame:
            txt = "\n".join([f"  P(lvl={p.level}): " + "".join([f"[{r.text}|font={r.font.name}|size={r.font.size.pt if r.font.size else None}|bold={r.font.bold}]" for r in p.runs]) for p in s.text_frame.paragraphs if p.text.strip()])
        print(f"Shape '{s.name}' ({s.shape_type}) at ({s.left}, {s.top}, {s.width}, {s.height}):\n{txt}")
